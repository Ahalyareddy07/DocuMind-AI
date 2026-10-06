"""Text extraction (PDF/DOCX/TXT), cleaning, chunking and statistics."""
import re
from pathlib import Path


class DocumentError(Exception):
    """Raised for unreadable, corrupted or empty documents."""


def clean_text(text: str) -> str:
    text = text.replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _paginate(text: str, size: int = 3000):
    """DOCX/TXT have no real pages, so group paragraphs into ~3000 char pseudo-pages."""
    pages, buf = [], ""
    for para in text.split("\n"):
        if len(buf) + len(para) > size and buf:
            pages.append(buf)
            buf = ""
        buf += para + "\n"
    if buf.strip():
        pages.append(buf)
    return list(enumerate(pages, 1))


def extract_pages(path: Path):
    """Return [(page_number, cleaned_text), ...]."""
    ext = Path(path).suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise DocumentError("This PDF is password protected.")
            raw = [(i, p.extract_text() or "") for i, p in enumerate(reader.pages, 1)]
        elif ext == ".docx":
            from docx import Document
            doc = Document(str(path))
            lines = [p.text for p in doc.paragraphs]
            for t in doc.tables:
                lines += [" | ".join(c.text for c in r.cells) for r in t.rows]
            raw = _paginate("\n".join(lines))
        else:
            raw = _paginate(Path(path).read_bytes().decode("utf-8", errors="replace"))
    except DocumentError:
        raise
    except Exception as e:  # corrupted file, parser failure, etc.
        raise DocumentError("The document appears to be corrupted or unreadable.") from e
    pages = [(n, clean_text(t)) for n, t in raw if clean_text(t)]
    if not pages:
        raise DocumentError("No text found. The document is empty or a scanned image (OCR is not supported).")
    return pages


def chunk_pages(doc_id: str, doc_name: str, pages, size: int = 900, overlap: int = 150):
    """Split each page into overlapping chunks, preferring sentence/word boundaries."""
    chunks = []
    for page, text in pages:
        start = 0
        while start < len(text):
            end = min(start + size, len(text))
            if end < len(text):
                cut = max(text.rfind(". ", start + size // 2, end), text.rfind("\n", start + size // 2, end))
                if cut == -1:
                    cut = text.rfind(" ", start + size // 2, end)
                end = cut + 1 if cut != -1 else end
            piece = text[start:end].strip()
            if piece:
                chunks.append({"chunk_id": f"{doc_id}-{len(chunks)}", "doc_id": doc_id,
                               "doc_name": doc_name, "page": page, "text": piece})
            if end >= len(text):
                break
            start = max(end - overlap, start + 1)
    return chunks


def compute_stats(pages) -> dict:
    text = "\n".join(t for _, t in pages)
    words = len(text.split())
    return {"pages": len(pages), "words": words, "characters": len(text),
            "paragraphs": len([p for p in text.split("\n") if p.strip()]),
            "reading_minutes": max(1, round(words / 200))}
